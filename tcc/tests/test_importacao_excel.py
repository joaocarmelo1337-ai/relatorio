import tempfile
import unittest
from datetime import date
from pathlib import Path

from tcc.database import db, seed
from tcc.services import importacao_excel as imp

try:
    import openpyxl
    TEM_OPENPYXL = True
except ImportError:
    TEM_OPENPYXL = False


def planilha_de_teste(caminho, obras, lancamentos):
    """Monta uma planilha com a mesma forma da do TCC, so que preenchida."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Obras")
    ws.append(["REGISTRO DAS UNIDADES INSPECIONADAS"])
    ws.append(["nota metodologica"])
    ws.append([])
    ws.append(["ID_Obra", "Identificação", "Endereço", "Data do habite-se",
               "Data de protocolo do projeto", "Regime normativo aplicável",
               "Data da vistoria", "Idade na vistoria (anos)", "Manual de uso?",
               "Plano de manutenção comprovado?", "Houve reforma?", "Observações"])
    for obra in obras:
        ws.append(obra)

    ws = wb.create_sheet("Catalogo")
    for _ in range(3):
        ws.append([])
    ws.append(["ID_Item", "Sistema / elemento", "Item de verificação",
               "Origem esperada", "Causas prováveis", "Procedência",
               "Tipo de falha", "Prazo de garantia (anos)", "Nota"])
    ws.append(["IT-009", "Estrutura", "Fissuras em vigas, pilares e lajes",
               "Endógena", "Retração do concreto", "Apêndice B",
               "Estrutura (Tabela 1)", 5, ""])
    ws.append(["IT-037", "Revestimento", "Falhas de rejunte", "Endógena",
               "Rejunte antes da cura", "Apêndice B",
               "Rejuntamento — desgaste", 1, ""])
    ws.append(["T3-04", "Vedações verticais", "Pinturas", "", "Lascamento",
               "NBR 17170 — Tabela 3", "Falha aparente", None, "Sem prazo em anos"])

    ws = wb.create_sheet("Lancamentos")
    for _ in range(3):
        ws.append([])
    ws.append(["ID_Lancamento", "ID_Obra", "ID_Item", "Data da vistoria",
               "Ocorrência (Sim/Não)", "Descrição", "G", "U", "T",
               "Pontuação GUT", "Prioridade", "Prazo", "Situação", "Observações"])
    for lancamento in lancamentos:
        ws.append(lancamento)

    wb.save(caminho)
    return caminho


@unittest.skipUnless(TEM_OPENPYXL, "openpyxl nao instalado")
class TestImportacao(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.con = db.criar_banco(self.dir / "teste.db")
        seed.semear_tudo(self.con)

    def tearDown(self):
        self.con.close()
        self.tmp.cleanup()

    def montar(self):
        obras = [
            ["UN-01", "Casa Silva", "Rua A, 100", date(2024, 3, 15),
             date(2024, 1, 10), None, date(2026, 9, 8), None, "Sim", "Não", "Não", ""],
            ["UN-02", "Casa Souza", "Rua B, 200", date(2021, 5, 20),
             date(2021, 2, 1), None, date(2026, 9, 8), None, "Não", "Não", "Sim", ""],
            ["UN-03", None, None, None, None, None, None, None, None, None, None, None],
        ]
        lancamentos = [
            ["LC-0001", "UN-01", "IT-009", date(2026, 9, 8), "Sim",
             "Fissura inclinada na viga V3", 8, 8, 8, None, None, None, None, "obs"],
            ["LC-0002", "UN-01", "IT-037", date(2026, 9, 8), "Sim",
             "Rejunte solto no box", 3, 3, 3, None, None, None, None, None],
            ["LC-0003", "UN-01", "IT-009", date(2026, 9, 8), "Não",
             "sem ocorrência", None, None, None, None, None, None, None, None],
            ["LC-0004", "UN-99", "IT-009", date(2026, 9, 8), "Sim",
             "obra inexistente", 6, 6, 6, None, None, None, None, None],
        ]
        return planilha_de_teste(self.dir / "planilha.xlsx", obras, lancamentos)

    def test_importa_obras_catalogo_e_lancamentos(self):
        resumo = imp.importar(self.con, self.montar())
        self.assertEqual(resumo["residencias"], 2)      # UN-03 esta em branco
        self.assertEqual(resumo["ocorrencias"], 2)      # LC-0003 e "Nao"; LC-0004 e orfa
        self.assertEqual(resumo["lancamentos_ignorados"], ["LC-0004"])

    def test_regime_normativo_e_calculado_na_importacao(self):
        imp.importar(self.con, self.montar())
        silva = self.con.execute(
            "SELECT regime_normativo FROM residencias WHERE codigo='UN-01'"
        ).fetchone()
        souza = self.con.execute(
            "SELECT regime_normativo FROM residencias WHERE codigo='UN-02'"
        ).fetchone()
        # protocolo 10/01/2024 -> depois de 10/06/2023 -> NBR 17170
        self.assertEqual(silva["regime_normativo"], "NBR 17170:2022")
        # protocolo 01/02/2021 -> regime anterior
        self.assertIn("Anexo D", souza["regime_normativo"])

    def test_gut_e_recalculado_e_nao_copiado(self):
        imp.importar(self.con, self.montar())
        linha = self.con.execute(
            "SELECT c.gut, c.prioridade FROM classificacoes_gut c "
            "JOIN ocorrencias o ON o.id = c.ocorrencia_id WHERE o.codigo='LC-0001'"
        ).fetchone()
        self.assertEqual(linha["gut"], 512)          # 8 x 8 x 8
        self.assertEqual(linha["prioridade"], 1)     # limite exato de P1

    def test_situacao_da_garantia_por_item(self):
        imp.importar(self.con, self.montar())
        # IT-009 (estrutura, 5 anos): Habite-se 15/03/2024 -> vence 15/03/2029, vigente
        estrutura = self.con.execute(
            "SELECT situacao_garantia FROM ocorrencias WHERE codigo='LC-0001'"
        ).fetchone()
        self.assertEqual(estrutura["situacao_garantia"], "VIGENTE")
        # IT-037 (rejuntamento, 1 ano): venceu em 15/03/2025
        rejunte = self.con.execute(
            "SELECT situacao_garantia FROM ocorrencias WHERE codigo='LC-0002'"
        ).fetchone()
        self.assertEqual(rejunte["situacao_garantia"], "VENCIDA")

    def test_a_data_da_vistoria_vira_vistoria(self):
        imp.importar(self.con, self.montar())
        n = self.con.execute("SELECT COUNT(*) n FROM vistorias").fetchone()["n"]
        self.assertEqual(n, 2)

    def test_reimportar_atualiza_em_vez_de_duplicar(self):
        caminho = self.montar()
        imp.importar(self.con, caminho)
        imp.importar(self.con, caminho)
        conta = lambda t: self.con.execute(f"SELECT COUNT(*) n FROM {t}").fetchone()["n"]
        self.assertEqual(conta("residencias"), 2)
        self.assertEqual(conta("ocorrencias"), 2)
        self.assertEqual(conta("vistorias"), 2)
        self.assertEqual(conta("classificacoes_gut"), 2)

    def test_recusa_planilha_sem_as_abas(self):
        caminho = self.dir / "vazia.xlsx"
        openpyxl.Workbook().save(caminho)
        with self.assertRaises(imp.PlanilhaInvalida):
            imp.importar(self.con, caminho)


if __name__ == "__main__":
    unittest.main()
